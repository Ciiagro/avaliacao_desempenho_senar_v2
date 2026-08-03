import io

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

AZUL_ESCURO = "16273F"
CINZA_CLARO = "EEF0F4"
BORDA_FINA = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def gerar_excel_resultado_final(dados):
    """Gera o Excel do Resultado Final, no mesmo layout da tabulação usada
    pela empresa, com fórmulas (SOMA, MÉDIA e Resultado Final recalculam se
    alguém editar uma nota)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultado final"

    fonte_titulo = Font(name="Arial", size=13, bold=True)
    fonte_subtitulo = Font(name="Arial", size=10, italic=True, color="666666")
    fonte_cabecalho = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fonte_normal = Font(name="Arial", size=10)
    fonte_negrito = Font(name="Arial", size=10, bold=True)
    fonte_resultado = Font(name="Arial", size=14, bold=True)
    fill_cabecalho = PatternFill("solid", fgColor=AZUL_ESCURO)
    fill_destaque = PatternFill("solid", fgColor=CINZA_CLARO)

    avaliado = dados["avaliado"]
    ciclo = dados["ciclo"]
    fatores = dados["fatores"]

    ws["A1"] = "Resultado final da avaliação de desempenho"
    ws["A1"].font = fonte_titulo
    ws["A2"] = f"{avaliado.nome} — Exercício {ciclo.exercicio}"
    ws["A2"].font = fonte_subtitulo

    col_widths = {"A": 5, "B": 46, "C": 18, "D": 22}
    for col, largura in col_widths.items():
        ws.column_dimensions[col].width = largura

    linha_cabecalho = 4
    cabecalhos = ["Nº", "FATOR", "(I) Autoavaliação", "(II) Avaliação do Gestor"]
    for i, texto in enumerate(cabecalhos):
        celula = ws.cell(row=linha_cabecalho, column=i + 1, value=texto)
        celula.font = fonte_cabecalho
        celula.fill = fill_cabecalho
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = BORDA_FINA

    primeira_linha_fator = linha_cabecalho + 1
    for i, fator in enumerate(fatores):
        linha = primeira_linha_fator + i
        nota_auto = dados["respostas_auto"].get(fator.id)
        nota_gestor = dados["respostas_gestor"].get(fator.id)

        ws.cell(row=linha, column=1, value=fator.ordem).font = fonte_normal
        ws.cell(row=linha, column=2, value=fator.nome).font = fonte_normal
        ws.cell(row=linha, column=3, value=nota_auto if nota_auto is not None else None).font = fonte_normal
        ws.cell(row=linha, column=4, value=nota_gestor if nota_gestor is not None else None).font = fonte_normal

        for col in range(1, 5):
            celula = ws.cell(row=linha, column=col)
            celula.border = BORDA_FINA
            if col in (1, 3, 4):
                celula.alignment = Alignment(horizontal="center")

    ultima_linha_fator = primeira_linha_fator + len(fatores) - 1

    linha_soma = ultima_linha_fator + 1
    ws.cell(row=linha_soma, column=2, value="SOMA").font = fonte_negrito
    if fatores:
        ws.cell(
            row=linha_soma, column=3,
            value=f"=SUM(C{primeira_linha_fator}:C{ultima_linha_fator})",
        )
        ws.cell(
            row=linha_soma, column=4,
            value=f"=SUM(D{primeira_linha_fator}:D{ultima_linha_fator})",
        )
    for col in range(1, 5):
        celula = ws.cell(row=linha_soma, column=col)
        celula.font = fonte_negrito
        celula.fill = fill_destaque
        celula.border = BORDA_FINA
        if col in (3, 4):
            celula.alignment = Alignment(horizontal="center")

    linha_media = linha_soma + 1
    qtd_fatores = len(fatores) or 1
    ws.cell(row=linha_media, column=2, value=f"MÉDIA (SOMA/{qtd_fatores})").font = fonte_negrito
    if fatores:
        ws.cell(row=linha_media, column=3, value=f"=C{linha_soma}/{qtd_fatores}").number_format = "0.00"
        ws.cell(row=linha_media, column=4, value=f"=D{linha_soma}/{qtd_fatores}").number_format = "0.00"
    for col in range(1, 5):
        celula = ws.cell(row=linha_media, column=col)
        celula.font = fonte_negrito
        celula.fill = fill_destaque
        celula.border = BORDA_FINA
        if col in (3, 4):
            celula.alignment = Alignment(horizontal="center")

    media_auto_ref = f"C{linha_media}"
    media_gestor_ref = f"D{linha_media}"

    linha_resultado = linha_media + 2
    ws.cell(row=linha_resultado, column=1, value="RESULTADO FINAL DA AVALIAÇÃO").font = fonte_negrito
    ws.merge_cells(start_row=linha_resultado, start_column=1, end_row=linha_resultado, end_column=2)
    ws.cell(row=linha_resultado, column=3, value="I x 0,30 + II x 0,70 =").font = fonte_normal
    celula_resultado = ws.cell(
        row=linha_resultado, column=4, value=f"={media_auto_ref}*0.3+{media_gestor_ref}*0.7"
    )
    celula_resultado.font = fonte_resultado
    celula_resultado.number_format = "0.00"
    celula_resultado.alignment = Alignment(horizontal="center")
    for col in range(1, 5):
        ws.cell(row=linha_resultado, column=col).border = BORDA_FINA

    resultado_ref = f"D{linha_resultado}"

    linha_conceito_topo = linha_resultado + 2
    ws.cell(row=linha_conceito_topo, column=1, value="CONCEITO").font = fonte_negrito
    ws.merge_cells(
        start_row=linha_conceito_topo, start_column=1, end_row=linha_conceito_topo + 2, end_column=1
    )
    ws.cell(row=linha_conceito_topo, column=1).alignment = Alignment(vertical="center")

    faixas = [
        ("A", "entre 10,0 e 8,0 (inclusive)", f'{resultado_ref}>=8'),
        ("B", "entre 7,9 e 7,0 (inclusive)", f'AND({resultado_ref}<8,{resultado_ref}>=7)'),
        ("C", "abaixo de 7,0", f'{resultado_ref}<7'),
    ]
    for i, (letra, descricao, condicao) in enumerate(faixas):
        linha = linha_conceito_topo + i
        marca = ws.cell(row=linha, column=2, value=f'=IF({condicao},"(X)","(  )")')
        marca.font = fonte_normal
        ws.cell(row=linha, column=3, value=f"{letra} — {descricao}").font = fonte_normal
        ws.merge_cells(start_row=linha, start_column=3, end_row=linha, end_column=4)
        for col in range(1, 5):
            ws.cell(row=linha, column=col).border = BORDA_FINA

    linha_local_data = linha_conceito_topo + 4
    ws.cell(row=linha_local_data, column=1, value="Local e data:").font = fonte_normal

    linha_avaliador = linha_local_data + 2
    avaliacao_gestor = dados["avaliacao_gestor"]
    avaliacao_auto = dados["avaliacao_auto"]
    nome_avaliador = avaliacao_gestor.avaliador.nome if avaliacao_gestor else "-"
    nome_avaliado_assinatura = avaliacao_auto.avaliado.nome if avaliacao_auto else avaliado.nome

    ws.cell(row=linha_avaliador, column=1, value="Assinatura do Avaliador:").font = fonte_normal
    ws.cell(row=linha_avaliador, column=2, value=nome_avaliador).font = fonte_normal
    ws.cell(row=linha_avaliador + 1, column=1, value="Assinatura do Avaliado:").font = fonte_normal
    ws.cell(row=linha_avaliador + 1, column=2, value=nome_avaliado_assinatura).font = fonte_normal

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# Metas de referência usadas na planilha da empresa (percentual esperado de
# cada conceito). Ajuste aqui se a meta mudar de um exercício para outro.
META_CONCEITO_A = 0.20
META_CONCEITO_B = 0.65
META_CONCEITO_C = 0.15


def gerar_excel_resultado_final_lista(ciclo, linhas, somente_concluidos=True):
    """Gera a planilha consolidada de todos os funcionários do ciclo, no
    mesmo formato da tabulação usada pela empresa: MAT, NOME, ADMISSÃO,
    NOTA e CONCEITO, com a tabela de referência x real dos conceitos ao
    final.

    `linhas` é a mesma lista montada na rota `admin.resultados` (uma
    dict por funcionário, com as chaves funcionario/media_auto/
    media_gestor/resultado_final/conceito/pronto).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultado final"

    fonte_titulo = Font(name="Arial", size=14, bold=True)
    fonte_subtitulo = Font(name="Arial", size=12, bold=True)
    fonte_cabecalho = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fonte_normal = Font(name="Arial", size=10)
    fonte_negrito = Font(name="Arial", size=10, bold=True)
    fill_cabecalho = PatternFill("solid", fgColor=AZUL_ESCURO)
    fill_destaque = PatternFill("solid", fgColor=CINZA_CLARO)
    centro = Alignment(horizontal="center", vertical="center")

    col_widths = {"A": 6, "B": 40, "C": 14, "D": 10, "E": 12}
    for col, largura in col_widths.items():
        ws.column_dimensions[col].width = largura

    ws.merge_cells("A1:E1")
    ws["A1"] = f"AVALIAÇÃO DE DESEMPENHO - Exercício {ciclo.exercicio}"
    ws["A1"].font = fonte_titulo
    ws["A1"].alignment = centro

    ws.merge_cells("A2:E2")
    ws["A2"] = "RESULTADO FINAL"
    ws["A2"].font = fonte_subtitulo
    ws["A2"].alignment = centro

    linha_cabecalho = 4
    cabecalhos = ["MAT", "NOME", "ADMISSÃO", "NOTA", "CONCEITO"]
    for i, texto in enumerate(cabecalhos):
        celula = ws.cell(row=linha_cabecalho, column=i + 1, value=texto)
        celula.font = fonte_cabecalho
        celula.fill = fill_cabecalho
        celula.alignment = centro

    linhas_a_exportar = [l for l in linhas if not somente_concluidos or l["pronto"]]

    contagem_conceitos = {"A": 0, "B": 0, "C": 0}
    primeira_linha_dado = linha_cabecalho + 1
    linha_atual = primeira_linha_dado
    for i, linha in enumerate(linhas_a_exportar, start=1):
        f = linha["funcionario"]
        ws.cell(row=linha_atual, column=1, value=f.matricula or i).alignment = centro
        ws.cell(row=linha_atual, column=2, value=f.nome).font = fonte_normal
        ws.cell(
            row=linha_atual, column=3,
            value=f.data_admissao.strftime("%d/%m/%Y") if f.data_admissao else "-",
        ).alignment = centro
        nota_cel = ws.cell(
            row=linha_atual, column=4,
            value=round(linha["resultado_final"], 2) if linha["resultado_final"] is not None else None,
        )
        nota_cel.alignment = centro
        nota_cel.number_format = "0.00"
        ws.cell(row=linha_atual, column=5, value=linha["conceito"] or "-").alignment = centro

        if linha["conceito"] in contagem_conceitos:
            contagem_conceitos[linha["conceito"]] += 1

        linha_atual += 1

    ultima_linha_dado = linha_atual - 1

    # Tabela de referência x real dos conceitos, como na planilha da empresa.
    total = sum(contagem_conceitos.values()) or 1
    linha_resumo = ultima_linha_dado + 3

    ws.merge_cells(start_row=linha_resumo, start_column=1, end_row=linha_resumo, end_column=2)
    ws.cell(row=linha_resumo, column=1, value="RESULTADO FINAL").font = fonte_negrito

    linha_resumo += 1
    for col, texto in enumerate(["CONCEITOS", "REFERÊNCIA", "REAL"]):
        celula = ws.cell(row=linha_resumo, column=col + 1, value=texto)
        celula.font = fonte_negrito
        celula.fill = fill_destaque
        celula.alignment = centro

    metas = {"A": META_CONCEITO_A, "B": META_CONCEITO_B, "C": META_CONCEITO_C}
    for letra in ("A", "B", "C"):
        linha_resumo += 1
        ws.cell(row=linha_resumo, column=1, value=f"CONCEITO {letra}").font = fonte_normal
        ws.cell(row=linha_resumo, column=2, value=metas[letra]).number_format = "0%"
        ws.cell(row=linha_resumo, column=3, value=contagem_conceitos[letra] / total).number_format = "0%"

    linha_resumo += 1
    ws.cell(row=linha_resumo, column=1, value="TOTAL").font = fonte_negrito
    ws.cell(row=linha_resumo, column=2, value=1).number_format = "0%"
    ws.cell(row=linha_resumo, column=3, value=1).number_format = "0%"
    for col in range(1, 4):
        ws.cell(row=linha_resumo, column=col).fill = fill_destaque
        ws.cell(row=linha_resumo, column=col).font = fonte_negrito

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
